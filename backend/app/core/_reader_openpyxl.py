"""`openpyxl` backend for `parser.py` (`XLSX_READER=openpyxl`).

This is the pre-calamine implementation, moved verbatim (only the imports of
shared helpers changed) — it exists so callers who need exact legacy
per-cell semantics (error cells, whitespace-only strings, dimension-only
trailing columns; see `_reader_calamine`'s module docstring for what
calamine cannot reproduce) can switch back to it.

`iter_chunks` streams a sheet in bounded-size batches via openpyxl read_only
mode, so a large primary file never fully materializes in memory.
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .exceptions import TemplateInvalid
from .parser import DEFAULT_CHUNK_SIZE, ParsedSheet, _make_sheet, _normalize_header, _row_is_empty


def list_sheets(path: str | Path) -> list[str]:
    try:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    except (InvalidFileException, FileNotFoundError, OSError, KeyError, BadZipFile) as exc:
        raise TemplateInvalid(
            user_message="範本檔損毀或非 xlsx 格式",
            tech_detail=f"{type(exc).__name__}: {exc}",
            path=str(path),
        ) from exc
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def preview_rows(path: str | Path, sheet: str, *, start: int = 1, count: int = 30) -> list[list[object]]:
    try:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    except (InvalidFileException, FileNotFoundError, OSError, BadZipFile) as exc:
        raise TemplateInvalid(
            user_message="範本檔損毀或非 xlsx 格式",
            tech_detail=f"{type(exc).__name__}: {exc}",
            path=str(path),
        ) from exc

    try:
        if sheet not in wb.sheetnames:
            raise TemplateInvalid(
                user_message=f"工作表「{sheet}」不存在",
                tech_detail=f"available sheets: {wb.sheetnames}",
                sheet=sheet,
            )
        ws = wb[sheet]
        out: list[list[object]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i < start:
                continue
            out.append(list(row))
            if len(out) >= count:
                break
        return out
    finally:
        wb.close()


def iter_chunks(
    path: str | Path, sheet: str, header_row: int, *, chunk_size: int = DEFAULT_CHUNK_SIZE
) -> Iterator[ParsedSheet]:
    if chunk_size < 1:
        chunk_size = DEFAULT_CHUNK_SIZE
    if header_row < 1:
        raise TemplateInvalid(
            user_message="標頭列號必須 ≥ 1",
            tech_detail=f"header_row={header_row}",
        )

    try:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    except (InvalidFileException, FileNotFoundError, OSError, BadZipFile) as exc:
        raise TemplateInvalid(
            user_message="範本檔損毀或非 xlsx 格式",
            tech_detail=f"{type(exc).__name__}: {exc}",
            path=str(path),
        ) from exc

    try:
        if sheet not in wb.sheetnames:
            raise TemplateInvalid(
                user_message=f"工作表「{sheet}」不存在",
                tech_detail=f"available sheets: {wb.sheetnames}",
                sheet=sheet,
            )

        ws = wb[sheet]
        headers: list[str] | None = None
        buffer: list[list[object]] = []
        emitted = False
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i < header_row:
                continue
            if i == header_row:
                headers = [_normalize_header(cell) for cell in row]
                continue
            if _row_is_empty(row):
                continue
            buffer.append(list(row))
            if len(buffer) >= chunk_size:
                yield _make_sheet(headers, buffer)
                emitted = True
                buffer = []

        if headers is None:
            raise TemplateInvalid(
                user_message=f"工作表「{sheet}」沒有第 {header_row} 列",
                tech_detail="header_row beyond last row in sheet",
                sheet=sheet,
                header_row=header_row,
            )

        if buffer:
            yield _make_sheet(headers, buffer)
            emitted = True
        if not emitted:
            yield _make_sheet(headers, [])
    finally:
        wb.close()
