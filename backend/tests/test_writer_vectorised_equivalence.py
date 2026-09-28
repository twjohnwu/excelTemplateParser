"""Equivalence tests for `writer._chunk_rows`, the vectorised replacement for the
per-row `chunk.iterrows()` + per-cell `row[col_name]` loop that used to build the
list `ws.append()`ed in `_write_fast`.

`ref_write_rows` below is the OLD loop body, copied verbatim (as `[list[list]]`
per chunk) so this test keeps asserting against the pre-change behaviour even
though the loop itself is gone from `writer.py`.
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from app.core import writer
from app.core.writer import _chunk_rows, _read_template_head


def ref_write_rows(chunk: pd.DataFrame, columns: list[str]) -> list[list[object]]:
    """The exact per-row list the OLD `_write_fast` loop passed to `ws.append()`."""
    out = []
    for _, row in chunk.iterrows():
        out_row = []
        for col_name in columns:
            value = row[col_name]
            if pd.isna(value):
                value = None
            out_row.append(value)
        out.append(out_row)
    return out


def assert_rows_equal(got: list[list[object]], ref: list[list[object]]) -> None:
    assert len(got) == len(ref)
    for row_got, row_ref in zip(got, ref):
        assert len(row_got) == len(row_ref)
        for v_got, v_ref in zip(row_got, row_ref):
            if v_ref is None:
                assert v_got is None
                continue
            assert type(v_got) is type(v_ref), (v_got, type(v_got), v_ref, type(v_ref))
            if isinstance(v_ref, float) and pd.isna(v_ref):
                assert pd.isna(v_got)
            else:
                assert v_got == v_ref


# ---------- dtype coverage (each case: `_chunk_rows` vs the old per-row loop) ----------


def test_nan_in_float_column():
    chunk = pd.DataFrame({"a": ["x", "y"], "金額": [1.5, np.nan]})
    columns = list(chunk.columns)
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


def test_none_in_object_column():
    chunk = pd.DataFrame({"a": ["x", None], "金額": [1, 2]})
    columns = list(chunk.columns)
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


def test_nat_in_datetime_column():
    chunk = pd.DataFrame({
        "a": ["x", "y"],
        "日期": pd.to_datetime(["2024-01-01", None]),
    })
    columns = list(chunk.columns)
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


def test_ints():
    chunk = pd.DataFrame({"a": ["x", "y", "z"], "數量": [1, 2, 3]})
    columns = list(chunk.columns)
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


def test_whole_number_floats():
    chunk = pd.DataFrame({"a": ["x", "y"], "金額": [10.0, 20.0]})
    columns = list(chunk.columns)
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


def test_bools():
    chunk = pd.DataFrame({"a": ["x", "y"], "啟用": [True, False]})
    columns = list(chunk.columns)
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


def test_strings():
    chunk = pd.DataFrame({"a": ["hello", "世界"], "b": ["foo", "bar"]})
    columns = list(chunk.columns)
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


def test_timestamp_value():
    chunk = pd.DataFrame({"a": ["x"], "日期": [pd.Timestamp("2021-05-05 10:30:00")]})
    columns = list(chunk.columns)
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


def test_realistic_mixed_row_types():
    """Product-pricing-shaped chunk: str + int + float(NaN) + datetime + bool
    together, the common production case (forces the object row dtype)."""
    chunk = pd.DataFrame({
        "SKU": ["SKU001", "SKU002", "SKU003"],
        "建議售價": [599, 1290, 690],
        "A報價": [394.0, np.nan, 463.0],
        "日期": pd.to_datetime(["2024-01-01", "2024-01-02", None]),
        "啟用": [True, False, True],
    })
    columns = list(chunk.columns)
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


def test_empty_chunk():
    chunk = pd.DataFrame({"a": pd.Series(dtype="object"), "b": pd.Series(dtype="float64")})
    columns = list(chunk.columns)
    assert _chunk_rows(chunk, columns) == ref_write_rows(chunk, columns) == []


def test_zero_length_dataframe_no_columns():
    chunk = pd.DataFrame()
    assert _chunk_rows(chunk, []) == ref_write_rows(chunk, []) == []


def test_column_order_different_from_target_order():
    """`_write_fast` looks values up by name (`row[col_name]`), so the incoming
    chunk's physical column order must not matter — only `columns` (the target
    order) does."""
    chunk = pd.DataFrame({"金額": [10, 20], "a": ["x", "y"]})  # reversed vs target
    columns = ["a", "金額"]  # target order
    assert_rows_equal(_chunk_rows(chunk, columns), ref_write_rows(chunk, columns))


# ---------- end-to-end: examples 01/02 through the real template ----------

EXAMPLES_DIR = Path(__file__).resolve().parents[2] / "examples"


def _write_fast_ref(src: Path, dst: Path, columns: list[str], chunks, sheet: str, header_row: int) -> None:
    """The pre-change `_write_fast`: identical structure, but the data-row loop
    still uses `chunk.iterrows()` + `ref_write_rows` instead of `_chunk_rows`."""
    head_rows, template_headers, col_widths = _read_template_head(src, sheet, header_row)

    name_to_pos: dict[str, int] = {}
    for pos, name in enumerate(template_headers):
        if name and name not in name_to_pos:
            name_to_pos[name] = pos
    new_cols = [c for c in columns if c not in name_to_pos]
    tmpl_width = len(template_headers)
    new_pos = {name: tmpl_width + k for k, name in enumerate(new_cols)}
    out_width = tmpl_width + len(new_cols)

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet)
    for letter, width in col_widths.items():
        ws.column_dimensions[letter].width = width

    for i in range(1, header_row + 1):
        src_row = head_rows[i - 1] if i - 1 < len(head_rows) else []
        out_row: list[object] = [None] * out_width
        for j, val in enumerate(src_row[:tmpl_width]):
            out_row[j] = val
        if i == header_row:
            for k, name in enumerate(new_cols):
                out_row[tmpl_width + k] = name
        ws.append(out_row)

    for chunk in chunks:
        for out_row_values in ref_write_rows(chunk, columns):
            out_row = [None] * out_width
            for col_name, value in zip(columns, out_row_values):
                pos = name_to_pos.get(col_name)
                if pos is None:
                    pos = new_pos[col_name]
                out_row[pos] = value
            ws.append(out_row)

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(dst))


def _example_dataframe(example: str) -> pd.DataFrame:
    return pd.read_excel(EXAMPLES_DIR / example / "expected_output.xlsx")


def _example_target(example: str) -> tuple[Path, str, int]:
    import json

    config = json.loads((EXAMPLES_DIR / example / "config.json").read_text(encoding="utf-8"))
    tmpl = config["target_template"]
    return EXAMPLES_DIR / example / "template.xlsx", tmpl["sheet"], tmpl["header_row"]


def _all_cell_values(path: Path, sheet: str) -> list[list[object]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _compare_example_end_to_end(example: str, tmp_path: Path) -> None:
    df = _example_dataframe(example)
    src, sheet, header_row = _example_target(example)
    columns = list(df.columns)

    new_out = tmp_path / f"{example}_new.xlsx"
    ref_out = tmp_path / f"{example}_ref.xlsx"

    writer.write_stream(src, [df], sheet=sheet, header_row=header_row, out_path=new_out, preserve_styles=False)
    _write_fast_ref(src, ref_out, columns, [df], sheet, header_row)

    assert _all_cell_values(new_out, sheet) == _all_cell_values(ref_out, sheet)


def test_example_01_product_pricing_end_to_end(tmp_path):
    _compare_example_end_to_end("01_product_pricing", tmp_path)


def test_example_02_agri_market_report_end_to_end(tmp_path):
    _compare_example_end_to_end("02_agri_market_report", tmp_path)
